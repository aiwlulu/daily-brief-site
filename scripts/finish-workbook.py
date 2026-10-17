"""Add the 重點 and 摘要明細 reading sheets and shorten edition links."""
import argparse
from copy import copy
import json
import math
import sys
import unicodedata

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = 'Noto Sans CJK TC'
INK, MUTED, LINE = '203C34', '5E7268', 'DCE4DE'
WHITE = PatternFill('solid', fgColor='FFFFFF')


def lines(text, size, width):
    # Estimate wrapped lines; Excel does not autofit row heights on open.
    total = 0
    for part in text.split('\n'):
        units = sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1.05 for c in part)
        total += max(1, math.ceil(units * size / 11 / (width * 0.95)))
    return total


def header(ws, data, width_cols):
    ws.sheet_view.showGridLines = False
    title = ws.cell(1, 1, data['title'])
    title.font = Font(name=FONT, size=16, bold=True, color=INK)
    ws.row_dimensions[1].height = 32
    note = ws.cell(2, 1, data['note'])
    note.font = Font(name=FONT, size=10, color=MUTED)
    note.alignment = Alignment(wrap_text=True, vertical='top')
    if width_cols > 1:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
    ws.row_dimensions[2].height = 34


def highlights(wb, data):
    # One row per story: number in A, bold headline and one-line detail in B.
    ws = wb.create_sheet('重點', 0)
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 96
    header(ws, data, 2)
    row = 4
    for i, entry in enumerate(data['entries'], 1):
        number = ws.cell(row, 1, i)
        number.font = Font(name=FONT, size=12, bold=True, color=MUTED)
        number.alignment = Alignment(horizontal='right', vertical='top')
        cell = ws.cell(row, 2)
        cell.value = CellRichText(
            TextBlock(InlineFont(rFont=FONT, sz=13, b=True, color=INK), entry['heading'] + '　' + entry['time'] + '\n'),
            TextBlock(InlineFont(rFont=FONT, sz=11, color=INK), entry['detail']),
        )
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        height = lines(entry['heading'] + '　' + entry['time'], 13, 96) * 20 + lines(entry['detail'], 11, 96) * 16 + 12
        ws.row_dimensions[row].height = height
        for col in (1, 2):
            ws.cell(row, col).border = Border(bottom=Side(style='thin', color=LINE))
            ws.cell(row, col).fill = WHITE
        row += 1
    ws.freeze_panes = 'A4'
    ws.sheet_properties.tabColor = 'C9A227'
    return ws


def details(wb, data):
    # Supplementary reading: bold headline row, full paragraph row below.
    ws = wb.create_sheet('摘要明細', 1)
    ws.column_dimensions['A'].width = 110
    header(ws, data, 1)
    row = 4
    for entry in data['entries']:
        heading = ws.cell(row, 1, entry['heading'] + '　' + entry['time'])
        heading.font = Font(name=FONT, size=13, bold=True, color=INK)
        heading.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row].height = lines(heading.value, 13, 110) * 20 + 6
        body = ws.cell(row + 1, 1, entry['paragraph'])
        body.font = Font(name=FONT, size=11, color=INK)
        body.alignment = Alignment(wrap_text=True, vertical='top')
        body.border = Border(bottom=Side(style='thin', color=LINE))
        ws.row_dimensions[row + 1].height = lines(entry['paragraph'], 11, 110) * 16 + 12
        row += 3
    ws.sheet_properties.tabColor = 'D9DFD6'


def short_links(wb):
    # Long URLs become a short clickable label; the table header is unchanged.
    ws = wb['市場動向']
    col = next(c.column for c in ws[1] if c.value == '日報連結')
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        if isinstance(cell.value, str) and cell.value.startswith('https://'):
            url = cell.value
            cell.value = '開啟週報' if 'view=week' in url else '開啟日報'
            cell.hyperlink = url
            font = copy(cell.font)
            font.color, font.underline = '1F6B56', 'single'
            cell.font = font


def build(path, data):
    wb = load_workbook(path)
    for name in ('重點', '摘要明細', '摘要'):
        if name in wb.sheetnames:
            del wb[name]
    details(wb, data)
    highlights(wb, data)
    short_links(wb)
    wb.active = 0
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet.title == '重點'
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--file', required=True)
    args = parser.parse_args()
    build(args.file, json.load(sys.stdin))
    print('Finished reading sheets in ' + args.file)


if __name__ == '__main__':
    main()
